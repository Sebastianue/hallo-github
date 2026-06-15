# -*- coding: utf-8 -*-
"""Erzeugt WM-2026-Prognose.xlsx aus wm_data.py."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import wm_data as D

wb = Workbook()

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
GRP_FILL = PatternFill("solid", fgColor="D9E1F2")
GRP_FONT = Font(bold=True, size=11, color="1F4E78")
PLAYED_FILL = PatternFill("solid", fgColor="E2EFDA")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER


# ---- Blatt 1: Gruppenphase ----
ws = wb.active
ws.title = "Gruppenphase"
ws["A1"] = "WM 2026 – Prognose der Gruppenspiele (Tipp vs. echtes Ergebnis)"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:J1")
ws["A2"] = f"Stand {D.STAND} · Fokus: Sieger- & Halbzeit-Tipp + Wertigkeit (Tendenz). Exaktes Ergebnis nur als Nebenangabe."
ws["A2"].font = Font(italic=True, size=9, color="808080")
ws.merge_cells("A2:J2")

headers = ["Gruppe", "Datum", "ST", "Begegnung", "Sieger-Tipp", "HZ-Tipp (Führung)",
           "Wertigkeit", "Endstand (echt)", "Treffer", "Erg.-Tipp (unsicher)"]
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, 10)

r = hr + 1
for (g, datum, st, home, away, sieger, erg, wert,
     hz_lead, hz_score, hz_wert, played, endstand, sok, hok) in D.GROUP_MATCHES_FULL:
    ws.cell(row=r, column=1, value=g).alignment = CENTER
    ws.cell(row=r, column=2, value=datum).alignment = CENTER
    ws.cell(row=r, column=3, value=st).alignment = CENTER
    ws.cell(row=r, column=4, value=f"{home} - {away}").alignment = LEFT
    ws.cell(row=r, column=5, value=sieger).alignment = LEFT
    ws.cell(row=r, column=6, value=(hz_lead if hz_lead == "—" else f"{hz_lead} ({hz_score})")).alignment = LEFT
    wc = ws.cell(row=r, column=7)
    if wert is None:
        wc.value = "—"
    else:
        wc.value = wert / 100.0
        wc.number_format = "0 %"
    wc.alignment = CENTER
    ws.cell(row=r, column=8, value=(endstand if played else "offen")).alignment = CENTER
    treffer = "offen" if not played else ("–" if sok is None else ("✓" if sok else "✗"))
    tcell = ws.cell(row=r, column=9, value=treffer)
    tcell.alignment = CENTER
    ws.cell(row=r, column=10, value=erg).alignment = CENTER
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        if played:
            cell.fill = PLAYED_FILL
    if played and sok is True:
        tcell.fill = OK_FILL
    elif played and sok is False:
        tcell.fill = BAD_FILL
    r += 1

for col, w in zip("ABCDEFGHIJ", [8, 12, 5, 28, 16, 18, 11, 14, 9, 18]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{hr}:J{hr + len(D.GROUP_MATCHES_FULL)}"

# ---- Blatt 2: Gruppentabellen-Prognose ----
ws2 = wb.create_sheet("Gruppentabellen-Prognose")
ws2["A1"] = "Prognostizierte Endplatzierungen je Gruppe"
ws2["A1"].font = TITLE_FONT
ws2.merge_cells("A1:E1")
for i, h in enumerate(["Gruppe", "1. Platz", "2. Platz", "3. Platz", "4. Platz (aus)"], 1):
    ws2.cell(row=3, column=i, value=h)
style_header(ws2, 3, 5)
r = 4
for row in D.STANDINGS:
    for i, val in enumerate(row, 1):
        cell = ws2.cell(row=r, column=i, value=val)
        cell.border = BORDER
        cell.alignment = CENTER if i == 1 else LEFT
    ws2.cell(row=r, column=1).fill = GRP_FILL
    ws2.cell(row=r, column=1).font = GRP_FONT
    r += 1
for col, w in zip("ABCDE", [8, 16, 16, 16, 16]):
    ws2.column_dimensions[col].width = w
ws2.auto_filter.ref = f"A3:E{3 + len(D.STANDINGS)}"

# ---- Blatt 3: K.-o.-Prognose ----
ws3 = wb.create_sheet("K.-o.-Prognose")
ws3["A1"] = "K.-o.-Phase – Turnierverlaufs-Prognose"
ws3["A1"].font = TITLE_FONT
ws3.merge_cells("A1:D1")
for i, h in enumerate(["Runde", "Datum", "Prognose", "Wertigkeit"], 1):
    ws3.cell(row=3, column=i, value=h)
style_header(ws3, 3, 4)
r = 4
for runde, datum, prog, wert in D.KO:
    ws3.cell(row=r, column=1, value=runde).alignment = LEFT
    ws3.cell(row=r, column=2, value=datum).alignment = CENTER
    ws3.cell(row=r, column=3, value=prog).alignment = LEFT
    wc = ws3.cell(row=r, column=4, value=wert / 100.0)
    wc.number_format = "0 %"
    wc.alignment = CENTER
    for c in range(1, 5):
        ws3.cell(row=r, column=c).border = BORDER
    r += 1
ws3.cell(row=r + 1, column=1, value="Titel-Favoriten-Ranking").font = GRP_FONT
r += 2
for i, h in enumerate(["Rang", "Team", "Titelchance"], 1):
    ws3.cell(row=r, column=i, value=h)
style_header(ws3, r, 3)
r += 1
for rang, team, chance in D.RANKING:
    ws3.cell(row=r, column=1, value=rang).alignment = CENTER
    ws3.cell(row=r, column=2, value=team).alignment = LEFT
    cc = ws3.cell(row=r, column=3, value=chance / 100.0)
    cc.number_format = "0 %"
    cc.alignment = CENTER
    for c in range(1, 4):
        ws3.cell(row=r, column=c).border = BORDER
    r += 1
for col, w in zip("ABCD", [22, 18, 36, 14]):
    ws3.column_dimensions[col].width = w

# ---- Blatt 4: Bilanz / Trefferquote ----
ws4 = wb.create_sheet("Bilanz")
ws4["A1"] = "Bilanz – meine bisherige Trefferquote (ehrlich)"
ws4["A1"].font = TITLE_FONT
ws4.merge_cells("A1:F1")
sk, sg, hk, hg, ex, eg = D.review_summary()
ws4["A2"] = (f"Sieger korrekt: {sk}/{sg} (~{round(100*sk/sg)} %) · "
             f"Halbzeit-Führung korrekt: {hk}/{hg} (~{round(100*hk/hg)} %) · "
             f"Exaktes Ergebnis: {ex}/{eg} (~{round(100*ex/eg)} %)")
ws4["A2"].font = Font(italic=True, size=10, color="404040")
ws4.merge_cells("A2:F2")
for i, h in enumerate(["Datum", "Begegnung", "Sieger-Tipp", "Endstand (echt)", "Sieger", "HZ"], 1):
    ws4.cell(row=4, column=i, value=h)
style_header(ws4, 4, 6)
r = 5
for (g, datum, st, home, away, sieger, erg, wert,
     hz_lead, hz_score, hz_wert, played, endstand, sok, hok) in D.GROUP_MATCHES_FULL:
    if not played:
        continue
    ws4.cell(row=r, column=1, value=datum[:6]).alignment = CENTER
    ws4.cell(row=r, column=2, value=f"{home} - {away}").alignment = LEFT
    ws4.cell(row=r, column=3, value=(sieger if sieger == "—" else f"{sieger} {erg}")).alignment = LEFT
    ws4.cell(row=r, column=4, value=endstand).alignment = CENTER
    ws4.cell(row=r, column=5, value=("–" if sok is None else ("✓" if sok else "✗"))).alignment = CENTER
    ws4.cell(row=r, column=6, value=("–" if hok is None else ("✓" if hok else "✗"))).alignment = CENTER
    for c in range(1, 7):
        ws4.cell(row=r, column=c).border = BORDER
    r += 1
for col, w in zip("ABCDEF", [9, 26, 18, 14, 8, 6]):
    ws4.column_dimensions[col].width = w
ws4.auto_filter.ref = f"A4:F{r - 1}"

wb.save("WM-2026-Prognose.xlsx")
print("WM-2026-Prognose.xlsx erstellt")
